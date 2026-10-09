#!/usr/bin/env python3
"""Read-only transformation of the study workbook and reviewed evidence seeds."""

from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "data/raw/study_plan.xlsx"
OUTPUT = ROOT / "public/data.json"


def clean(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if value == "":
        return None
    return value


def workbook_rows(sheet_name: str) -> list[list[Any]]:
    if not WORKBOOK.exists():
        raise FileNotFoundError(f"Missing read-only workbook copy: {WORKBOOK}")
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        return [[clean(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def nonempty_rows(sheet_name: str) -> list[list[Any]]:
    return [row for row in workbook_rows(sheet_name) if any(value is not None for value in row)]


def number(value: Any) -> int | float | None:
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def assign_district(unit: str | None, districts: list[dict[str, Any]]) -> tuple[str | None, str | None]:
    if not unit:
        return None, None
    matches = [
        (district["id"], term)
        for district in districts
        for term in [district["name"], *district.get("unitMatchAliases", [])]
        if term and term in unit
    ]
    matched_districts = {district_id for district_id, _ in matches}
    if len(matched_districts) != 1:
        return None, None
    district_id = next(iter(matched_districts))
    match_term = max(
        (term for matched_id, term in matches if matched_id == district_id),
        key=len,
    )
    return district_id, match_term


def parse_mirror_major_criteria(major_text: str | None) -> dict[str, list[str]]:
    text = str(major_text or "")
    criteria: dict[str, list[str]] = {}
    for level, label in (("undergraduate", "本科"), ("graduate", "研究生")):
        section = re.search(
            rf"{label}\s*[:：](.*?)(?=(?:本科|研究生)\s*[:：]|$)",
            text,
        )
        if not section:
            continue
        codes = re.findall(r"[（(]\s*([A-Za-z0-9]{2,})\s*[）)]", section.group(1))
        if codes:
            criteria[level] = list(dict.fromkeys(codes))
    return criteria


def build_positions(
    positions: list[dict[str, Any]], districts: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    result = []
    for position in positions:
        district_id, district_match = assign_district(position.get("unit"), districts)
        mirror_sources = position.get("sources", [])
        is_shijingshan_mirror = any(
            source_id.startswith("gwyzwb-2025-shijingshan-org-")
            or source_id == "huatu-2024-shijingshan-list"
            for source_id in mirror_sources
        )
        derived_criteria = (
            parse_mirror_major_criteria(position.get("majorText"))
            if is_shijingshan_mirror and not position.get("majorCriteria")
            else {}
        )
        result.append({
            **position,
            **({"majorCriteria": derived_criteria} if derived_criteria else {}),
            "districtId": district_id,
            "districtMatch": district_match,
        })
    return result


def build_study_plan() -> tuple[list[dict[str, Any]], dict[str, int | float]]:
    rows = nonempty_rows("50天计划")
    days: list[dict[str, Any]] = []
    for row in rows[1:]:
        day = number(row[0])
        if day is None:
            continue
        days.append({
            "day": int(day),
            "date": row[1],
            "stage": row[2],
            "focus": row[3],
            "coreTask": row[4],
            "plannedQuestions": number(row[5]),
            "actualQuestions": number(row[6]),
            "plannedHours": number(row[7]),
            "actualHours": number(row[8]),
            "status": row[9] or "未开始",
            "essayTask": row[11],
            "politicalTask": row[12],
            "reviewNote": row[13],
        })
    meta = {
        "days": len(days),
        "plannedQuestions": sum(day["plannedQuestions"] or 0 for day in days),
        "plannedHours": sum(day["plannedHours"] or 0 for day in days),
    }
    return days, meta


def build_aptitude() -> list[dict[str, Any]]:
    rows = nonempty_rows("行测清单")
    result = []
    for row in rows[1:]:
        if not row[0] or not row[1]:
            continue
        attempted = number(row[5])
        if attempted == 0:
            attempted = None
        result.append({
            "area": row[0],
            "item": row[1],
            "targetAccuracy": number(row[2]),
            "timeTarget": row[3],
            "plannedQuestions": number(row[4]),
            "attempted": attempted,
            "accuracy": number(row[6]) if attempted is not None and attempted > 0 else None,
            "retakeAccuracy": number(row[7]) if attempted is not None and attempted > 0 else None,
            "mastery": number(row[8]) if attempted is not None and attempted > 0 else None,
            "passStandard": row[9],
        })
    return result


def build_essay() -> list[dict[str, Any]]:
    rows = nonempty_rows("申论清单")
    result = []
    for row in rows[1:]:
        if not row[0] or not row[1]:
            continue
        completed = number(row[3])
        if completed == 0:
            completed = None
        result.append({
            "area": row[0],
            "practice": row[1],
            "planned": number(row[2]),
            "completed": completed,
            "selfScore": number(row[4]) if completed is not None and completed > 0 else None,
            "keywordCoverage": number(row[5]) if completed is not None and completed > 0 else None,
            "timedPass": bool(row[6]) if completed is not None and completed > 0 else None,
            "mastery": number(row[7]) if completed is not None and completed > 0 else None,
            "acceptance": row[8],
        })
    return result


def build_score_rows(sources: list[dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    sources = sources or load_json(ROOT / "data/source_registry.json")
    reviewed_rows = load_json(ROOT / "data/score_rows_seed.json")
    if reviewed_rows:
        source_by_id = {source["sourceId"]: source for source in sources}
        orphaned = [row.get("sourceId") for row in reviewed_rows if row.get("sourceId") not in source_by_id]
        if orphaned:
            raise ValueError(f"Reviewed cutoff row has no registered source: {orphaned[0]!r}")
        return [{
            **row,
            "sourceLevel": source_by_id[row["sourceId"]]["level"],
            "verification": row.get("verification") or (
                "年度、单位与岗位名唯一匹配"
                if row.get("mappingConfidence") == "high"
                else "岗位代码未唯一匹配；保留空代码"
            ),
        } for row in reviewed_rows]
    source_by_url = {source["url"]: source for source in sources}
    rows = nonempty_rows("2026参考数据")
    in_section = False
    result = []
    for row in rows:
        if row[0] == "昌平2026样本职位":
            in_section = True
            continue
        if row[0] in {"与你专业相关的2026昌平样本", "官方规则"}:
            break
        if not in_section or not row[0] or number(row[1]) is None:
            continue
        source = source_by_url.get(row[3])
        if not source:
            raise ValueError(f"No registered source for workbook score sample: {row[3]!r}")
        name = str(row[0])
        org_type = "街道" if "街道" in name else "镇" if "镇" in name else "未知"
        result.append({
            "year": 2026,
            "name": name,
            "positionCode": None,
            "orgType": org_type,
            "score": number(row[1]),
            "notes": row[2],
            "sourceId": source["sourceId"],
            "sourceLevel": source["level"],
            "verification": "第三方职位样例，未逐条独立核验",
        })
    return result


def build_mocks() -> list[dict[str, Any]]:
    rows = nonempty_rows("模考记录")
    result = []
    for row in rows[1:]:
        mock_no = number(row[0])
        if mock_no is None:
            continue
        aptitude = number(row[2])
        essay = number(row[3])
        result.append({
            "number": int(mock_no),
            "date": row[1],
            "aptitude": aptitude,
            "essay": essay,
            "total": aptitude + essay if aptitude is not None and essay is not None else None,
            "aptitudeMinutes": number(row[5]),
            "essayMinutes": number(row[6]),
            "accuracy": {
                "dataAnalysis": number(row[7]),
                "reasoning": number(row[8]),
                "science": number(row[9]),
                "quantitative": number(row[10]),
                "verbal": number(row[11]),
                "politicalAndGeneral": number(row[12]),
            },
            "target": number(row[13]),
            "review": row[15],
        })
    return result


def build_aggregate_evidence(sources: list[dict[str, Any]], score_rows: list[dict[str, Any]] | None = None) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    source_by_id = {source["sourceId"]: source for source in sources}
    observations = []
    for source_id, qualified, recruit_count, observed_at in [
        ("eoffcn-2024-qualified", 1671, 199, "2023-11-17 09:00"),
        ("eoffcn-2025-qualified", 1738, 176, "2024-11-22 09:00"),
    ]:
        source = source_by_id[source_id]
        observations.append({
            "year": source["year"],
            "scope": "昌平区报名快照（区级汇总，非岗位级）",
            "positionCode": None,
            "observationType": "qualified_snapshot",
            "observedAt": observed_at,
            "applicantsRegistered": None,
            "applicantsQualified": qualified,
            "applicantsPaid": None,
            "applicantsConfirmed": None,
            "actualTestTakers": None,
            "recruitCount": recruit_count,
            "sourceId": source_id,
            "sourceLevel": "secondary",
            "notes": source["notes"],
        })
    for source in sources:
        for snapshot in source.get("positionSnapshots", []):
            position_code = snapshot.get("positionCode")
            observations.append({
                "year": source["year"],
                "scope": snapshot.get("scope", "position-level" if position_code else "unit-level"),
                "positionCode": position_code,
                "unit": snapshot.get("unit"),
                "observationType": "qualified_snapshot",
                "observedAt": source["observedAt"],
                "applicantsRegistered": None,
                "applicantsQualified": snapshot.get("applicantsQualified"),
                "applicantsPaid": None,
                "applicantsConfirmed": None,
                "actualTestTakers": None,
                "recruitCount": snapshot.get("recruitCount"),
                "sourceId": source["sourceId"],
                "sourceLevel": source["level"],
                "notes": source["notes"],
            })

    ratio_source = source_by_id["sohu-2026-competition"]
    observations.append({
        "year": 2026,
        "scope": "昌平区平均竞争比报道（口径未详）",
        "positionCode": None,
        "observationType": "unknown_ratio_snapshot",
        "observedAt": "2025-11-21 09:00",
        "applicantsRegistered": None,
        "applicantsQualified": None,
        "applicantsPaid": None,
        "applicantsConfirmed": None,
        "actualTestTakers": None,
        "recruitCount": None,
        "unknownRegistrationRatio": ratio_source["reportedAverageRatio"],
        "sourceId": ratio_source["sourceId"],
        "sourceLevel": "secondary",
        "notes": ratio_source["notes"],
    })

    conflicts = []
    for year, source_ids in {
        2024: ["huatu-2024-list", "gwyzwb-2024-list", "gaodun-2024-list"],
        2025: ["huatu-2025-list", "gwyzwb-2025-list", "eoffcn-2025-list"],
        2026: ["huatu-2026-list", "gwyzwb-2026-list"],
    }.items():
        for metric, field in [("position_count", "reportedPositionCount"), ("recruit_count", "reportedRecruitCount")]:
            for source_id in source_ids:
                value = source_by_id[source_id].get(field)
                if value is not None:
                    conflicts.append({
                        "metric": metric,
                        "year": year,
                        "value": value,
                        "sourceId": source_id,
                        "level": "secondary",
                        "notes": source_by_id[source_id].get("notes", ""),
                    })

    score_rows = score_rows or []
    score_samples = []
    years = sorted({int(row["year"]) for row in score_rows if number(row.get("score")) is not None})
    for year in years:
        year_rows = [row for row in score_rows if int(row["year"]) == year and number(row.get("score")) is not None]
        source_ids = sorted({row["sourceId"] for row in year_rows})
        for source_id in source_ids:
            source_rows = [row for row in year_rows if row["sourceId"] == source_id]
            source = source_by_id[source_id]
            values = [number(row["score"]) for row in source_rows]
            score_samples.append({
                "year": year,
                "scope": "昌平区岗位最低进面线部分样本",
                "samplePositions": len(source_rows),
                "sampleRecruits": source.get("sampleRecruits"),
                "minimum": min(values),
                "maximum": max(values),
                "sourceId": source_id,
                "sourceLevel": source["level"],
                "complete": False,
                "notes": source["notes"],
            })
    return observations, conflicts, score_samples


def render_coverage_report(dataset: dict[str, Any]) -> str:
    positions = dataset.get("positions", [])
    score_rows = dataset.get("scoreRows", [])
    score_samples = dataset.get("scoreSamples", [])
    observations = dataset.get("observations", [])
    sources = dataset.get("sources", [])
    official_counts = dataset.get("coverage", {}).get("officialPositionCounts", {})

    def has_value(value: Any) -> bool:
        return value is not None and value != ""

    def link_source(source: dict[str, Any]) -> str:
        source_id = source.get("sourceId", "unknown")
        url = source.get("url")
        return f"[{source_id}]({url})" if url else str(source_id)

    lines = [
        "# 昌平京考看板数据覆盖报告",
        "",
        f"数据基准日：{dataset.get('dataAsOf', '未知')}。本报告由 `scripts/build_data.py` 根据生成数据自动更新。年度官方职位分母未取得时只展示可见样例和二手参考，不把参考数写成官方覆盖率。",
        "",
        "## 年度可见职位",
        "",
        "| 年度 | 已导入职位行 | 已知招录人数 | 具名分数（唯一代码关联 / 样例行） | 官方职位行 |",
        "| ---: | ---: | ---: | ---: | ---: |",
    ]

    for year in (2024, 2025, 2026):
        year_positions = [position for position in positions if position.get("year") == year]
        known_recruits = [number(position.get("recruitCount")) for position in year_positions]
        known_recruits = [count for count in known_recruits if count is not None]
        year_scores = [row for row in score_rows if row.get("year") == year]
        linked_scores = [
            row for row in year_scores
            if row.get("mappingConfidence") == "high" and row.get("positionCode")
        ]
        official_rows = sum(position.get("sourceLevel") == "official" for position in year_positions)
        lines.append(
            f"| {year} | {len(year_positions)} | {sum(known_recruits)} | "
            f"{len(linked_scores)} / {len(year_scores)} | {official_rows} |"
        )

    total_positions = len(positions)
    total_recruits = sum(number(position.get("recruitCount")) or 0 for position in positions)
    source_levels = {
        level: sum(source.get("level") == level for source in sources)
        for level in ("official", "secondary")
    }
    lines.extend([
        "",
        f"合计 {total_positions} 条候选职位、{total_recruits} 个已知招录名额；来源注册表 {len(sources)} 项（官方 {source_levels['official']}、二手 {source_levels['secondary']}）。来源条目数不等于逐岗官方核验数。",
        "",
        "## 结构化字段覆盖",
        "",
        "字段有值只表示当前镜像可见，不表示完整或已由官方核实。`专业代码目录` 不等同于资格条件完整。",
        "",
        "| 年度 | 专业代码目录 | 结构化要求 | 资格完整 | 二手交叉字段一致 | 职位简介 | 专业测试字段 | 体测字段 | 面试比例 | 电话 / 网站 |",
        "| ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: |",
    ])
    for year in (2024, 2025, 2026):
        year_positions = [position for position in positions if position.get("year") == year]
        count = len(year_positions)
        field_counts = [
            sum(bool(position.get("majorCriteria")) for position in year_positions),
            sum(bool(position.get("requirements")) for position in year_positions),
            sum(position.get("eligibilityComplete") is True for position in year_positions),
            sum(position.get("crossVerified") is True for position in year_positions),
            sum(bool(position.get("positionDescription")) for position in year_positions),
            sum(has_value(position.get("professionalTest")) for position in year_positions),
            sum(has_value(position.get("physicalTest")) for position in year_positions),
            sum(bool(position.get("interviewRatio")) for position in year_positions),
            sum(bool(position.get("consultPhone")) or bool(position.get("unitWebsite")) for position in year_positions),
        ]
        lines.append(f"| {year} | " + " | ".join(f"{value} / {count}" for value in field_counts) + " |")

    lines.extend([
        "",
        "## 二手年度汇总与职位差额",
        "",
        "| 年度 | 可见二手汇总（职位 / 招录） | 与当前行级样例的关系 |",
        "| ---: | --- | --- |",
    ])
    for year in (2024, 2025, 2026):
        references = [
            source for source in sources
            if source.get("year") == year
            and source.get("level") == "secondary"
            and source.get("geographicScope") != "citywide"
            and not source.get("districtId")
            and (has_value(source.get("reportedPositionCount")) or has_value(source.get("reportedRecruitCount")))
        ]
        reference_text = "; ".join(
            f"{link_source(source)} {source.get('reportedPositionCount', '未知')} 岗 / {source.get('reportedRecruitCount', '未知')} 人"
            for source in references
        ) or "未登记可比二手总数"
        year_positions = [position for position in positions if position.get("year") == year]
        known_recruits = [number(position.get("recruitCount")) for position in year_positions]
        known_recruits = [count for count in known_recruits if count is not None]
        totals = {
            (source.get("reportedPositionCount"), source.get("reportedRecruitCount"))
            for source in references
        }
        if len(totals) == 1 and references:
            reference_positions, reference_recruits = next(iter(totals))
            position_gap = max(0, reference_positions - len(year_positions)) if number(reference_positions) is not None else None
            recruit_gap = max(0, reference_recruits - sum(known_recruits)) if number(reference_recruits) is not None else None
            relation = (
                f"{len(year_positions)} / {reference_positions} 岗、{sum(known_recruits)} / {reference_recruits} 人；"
                f"相差 {position_gap if position_gap is not None else '未知'} 岗 / "
                f"{recruit_gap if recruit_gap is not None else '未知'} 人（仅二手相对比较，非官方覆盖率）"
            )
        elif references:
            relation = "二手来源总数不一致；不计算年度覆盖率或总量差额"
        else:
            relation = "无可比二手分母；年度覆盖率未知"
        lines.append(f"| {year} | {reference_text} | {relation} |")

    citywide_references = [
        source for source in sources
        if source.get("level") == "secondary"
        and source.get("geographicScope") == "citywide"
        and not source.get("districtId")
        and (has_value(source.get("reportedPositionCount")) or has_value(source.get("reportedRecruitCount")))
    ]
    if citywide_references:
        lines.extend([
            "",
            "## 全市级职位表参照汇总",
            "",
            "以下为全市范围的二手汇总，仅作外部参照；不是官方原表复算结果，不与昌平逐岗样本计算差额，也不作为官方覆盖率分母。",
            "",
            "| 年度 | 来源 | 全市职位 / 招录 | 说明 |",
            "| ---: | --- | ---: | --- |",
        ])
        for source in citywide_references:
            lines.append(
                f"| {source.get('year', '未知')} | {link_source(source)} | "
                f"{source.get('reportedPositionCount', '未知')} 岗 / "
                f"{source.get('reportedRecruitCount', '未知')} 人 | "
                "不与昌平逐岗样本计算差额 |"
            )

    unit_reference_source = next((
        source for source in sources
        if source.get("year") == 2026 and source.get("sourceId") == "huatu-2026-list"
        and source.get("positionUnitReferences")
    ), None)
    if unit_reference_source is None:
        unit_reference_source = next((
            source for source in sources
            if source.get("year") == 2026 and source.get("positionUnitReferences")
        ), None)
    unit_gap_intro = (
        f"单位汇总参考来源为 `{unit_reference_source.get('sourceId', 'unknown')}`；以下 2026 单位级差额仅用于定位二手镜像中可见的收录差异，不等于已取得缺失岗位代码或官方核验。"
        if unit_reference_source
        else "未登记 2026 单位级职位数汇总；当前无法逐单位列出二手镜像差额。"
    )
    lines.extend([
        "",
        unit_gap_intro,
        "",
        "| 单位 | 镜像岗位 / 招录 | 已导入岗位 / 招录 | 镜像差额 |",
        "| --- | ---: | ---: | ---: |",
    ])
    if unit_reference_source:
        for reference in unit_reference_source.get("positionUnitReferences", []):
            unit = reference.get("unit")
            unit_positions = [position for position in positions if position.get("year") == 2026 and position.get("unit") == unit]
            known_unit_recruits = [number(position.get("recruitCount")) for position in unit_positions]
            known_unit_recruits = [count for count in known_unit_recruits if count is not None]
            expected_positions = number(reference.get("positionCount"))
            expected_recruits = number(reference.get("recruitCount"))
            position_gap = max(0, expected_positions - len(unit_positions)) if expected_positions is not None else None
            recruit_gap = max(0, expected_recruits - sum(known_unit_recruits)) if expected_recruits is not None else None
            gap_text = (
                f"缺 {position_gap if position_gap is not None else '未知'} 岗 / "
                f"{recruit_gap if recruit_gap is not None else '未知'} 人"
            )
            lines.append(
                f"| {unit} | {expected_positions if expected_positions is not None else '未知'} / "
                f"{expected_recruits if expected_recruits is not None else '未知'} | "
                f"{len(unit_positions)} / {sum(known_unit_recruits)} | {gap_text} |"
            )

    scope_counts = {
        "position-level": sum(observation.get("scope") == "position-level" for observation in observations),
        "unit-level": sum(observation.get("scope") == "unit-level" for observation in observations),
    }
    district_scoped = sum(
        isinstance(observation.get("scope"), str) and observation["scope"].startswith("昌平区")
        for observation in observations
    )
    registered_counts = sum(has_value(observation.get("applicantsRegistered")) for observation in observations)
    qualified_counts = sum(has_value(observation.get("applicantsQualified")) for observation in observations)
    paid_counts = sum(has_value(observation.get("applicantsPaid")) for observation in observations)
    confirmed_counts = sum(has_value(observation.get("applicantsConfirmed")) for observation in observations)
    test_taker_counts = sum(has_value(observation.get("actualTestTakers")) for observation in observations)
    lines.extend([
        "",
        "## 进面线与报名观察",
        "",
        f"具名最低进面分共 {len(score_rows)} 行（唯一代码关联 {sum(row.get('mappingConfidence') == 'high' and bool(row.get('positionCode')) for row in score_rows)} 行）；分数样本范围记录 {len(score_samples)} 条。面试最低进面线不等同笔试合格线。",
        f"范围观察共 {len(observations)} 条：岗位级 {scope_counts['position-level']}、单位级 {scope_counts['unit-level']}、区级范围 {district_scoped}。其中有资格审查通过数的 {qualified_counts} 条；报名数 {registered_counts}、缴费数 {paid_counts}、确认参考数 {confirmed_counts}、实际参考人数 {test_taker_counts} 条。过程快照不等同最终人数。",
        "",
        "## 校验结论",
        "",
        "COVERAGE: INCOMPLETE — 年度官方职位总量分母未核实。结构校验请运行 `python3 scripts/validate_data.py`；本报告不把二手参考数表述为官方全量。",
        "",
        "2024 年公开二手汇总存在总量差异，2025 年汇总也有小幅差异；职位代码集合差集尚未完成。2024、2025 暂无具名进面分行。原始学习计划 Excel 未被修改。",
        "",
    ])
    return "\n".join(lines)


def main() -> None:
    sources = load_json(ROOT / "data/source_registry.json")
    districts = load_json(ROOT / "data/districts.json")
    positions = build_positions(load_json(ROOT / "data/positions_seed.json"), districts)
    study_plan, study_meta = build_study_plan()
    score_rows = build_score_rows(sources)
    observations, conflicts, score_samples = build_aggregate_evidence(sources, score_rows)
    output = {
        "schemaVersion": 1,
        "generatedAt": date.today().isoformat(),
        "dataAsOf": date.today().isoformat(),
        "scopeNote": "职位库仍为可追溯候选，不是北京全市全量官方职位表。现有行级候选覆盖昌平、延庆、石景山2024—2026年，以及海淀2025—2026年；石景山2024、海淀2025和海淀2026各保留一条仅见于来源目录的市级候选，因单位名不能确认区属而未分配区县。海淀2024目前只有第三方区级汇总，没有可追溯逐岗行，不把汇总数合成为职位。年度官方全市分母及逐代码核验尚未完成。区目录按北京市民政局2026年行政区划代码标准化；职位区县仅在招录单位名匹配标准区名或经审阅别名时赋值，其余保留未知。职位事实与报名快照、部分面试分数样本分开保存。",
        "studyMeta": study_meta,
        "districts": districts,
        "positions": positions,
        "observations": observations,
        "sources": sources,
        "studyPlan": study_plan,
        "aptitude": build_aptitude(),
        "essay": build_essay(),
        "mocks": build_mocks(),
        "conflicts": conflicts,
        "scoreSamples": score_samples,
        "scoreRows": score_rows,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    coverage_report = ROOT / "research/coverage_report.md"
    coverage_report.parent.mkdir(parents=True, exist_ok=True)
    coverage_report.write_text(render_coverage_report(output), encoding="utf-8")
    print(f"Built {OUTPUT.relative_to(ROOT)}: {len(positions)} candidate jobs, {len(study_plan)} study days, {len(sources)} sources")


if __name__ == "__main__":
    main()
